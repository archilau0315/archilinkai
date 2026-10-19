"""V0.59 新建+拆除双估算单测：拆除选表、split 分组、合计、多 sheet 导出。"""

from __future__ import annotations

import pytest

from archilinkai import aggregate_service
from archilinkai import instance_manager
from archilinkai import server
from archilinkai.cost_data import MACHINE_BOOK, PRICE_BOOK
from archilinkai.cost_estimator import (
    CONSUMPTION,
    DEMOLITION_CONSUMPTION,
    build_combined_estimate_report,
    combine_estimate_tables,
    estimate_from_flat_rows,
    split_rows_by_renovation,
)


def _rows(code="010502001", name="现浇混凝土矩形柱", unit="m³", qty=100.0,
          renovation=None, story=0, etype="Column"):
    row = {
        "story": story, "type": etype, "count": 1,
        "code": code, "code_name": name, "code_unit": unit, "sum": qty,
    }
    if renovation is not None:
        row["renovation"] = renovation
    return row


# --- 拆除消耗量表 ---------------------------------------------------------

def test_demolition_covers_all_consumption_codes():
    missing = set(CONSUMPTION) - set(DEMOLITION_CONSUMPTION)
    assert not missing, f"拆除表缺编码：{missing}"
    assert len(DEMOLITION_CONSUMPTION) >= 38


def test_demolition_labor_higher_than_build():
    """拆除更耗工：拆除人工高于同编码建造人工。"""
    lower = [c for c in CONSUMPTION
             if DEMOLITION_CONSUMPTION[c]["labor_qty"]
             <= CONSUMPTION[c]["labor_qty"]]
    assert not lower, f"拆除人工未高于建造的编码：{lower}"


def test_demolition_prices_and_machines_present():
    assert "渣土清运" in PRICE_BOOK
    assert "切割片" in PRICE_BOOK
    for name in ("破碎锤", "混凝土切割机", "自卸汽车"):
        assert name in MACHINE_BOOK


# --- 选表生效 -------------------------------------------------------------

def test_consumption_table_selects_demolition():
    rows = [_rows(qty=100.0)]
    build = estimate_from_flat_rows(rows, "P")
    demo = estimate_from_flat_rows(
        rows, "P", consumption_table=DEMOLITION_CONSUMPTION)
    b, d = build["items"][0], demo["items"][0]
    assert b["calc_mode"] == d["calc_mode"] == "detailed"
    assert d["labor_price"] != b["labor_price"]
    assert d["material_price"] != b["material_price"]
    assert d["machine_price"] != b["machine_price"]


def test_consumption_table_default_is_build():
    rows = [_rows(qty=100.0)]
    default = estimate_from_flat_rows(rows, "P")
    explicit = estimate_from_flat_rows(
        rows, "P", consumption_table=CONSUMPTION)
    assert default["summary"] == explicit["summary"]


def test_demolition_uncovered_code_notes():
    rows = [_rows(code="999999999", name="神秘工艺构件", qty=10.0)]
    demo = estimate_from_flat_rows(
        rows, "P", consumption_table=DEMOLITION_CONSUMPTION)
    assert demo["items"][0]["calc_mode"] == "skipped"
    assert demo["items"][0]["total_price"] is None
    assert any("无消耗量参考" in n for n in demo["notes"])


# --- split 纯函数 ---------------------------------------------------------

def test_split_rows_by_renovation_pure():
    rows = [_rows(renovation="New"), _rows(renovation="ToBeDemolished"),
            _rows(renovation="Existing"), _rows(renovation="Undefined")]
    new, demo, existing = split_rows_by_renovation(rows)
    assert len(new) == 1 and len(demo) == 1 and existing == 1


# --- estimate_project（调用 server 注册的工具函数）-----------------------

def _tool():
    return server.mcp._tool_manager._tools["estimate_project"].fn


class _Info:
    instance_id = "ac28-1"
    project_name = "办公楼改造"


_FLAT = [
    _rows(renovation="New", code="010502001", name="现浇混凝土矩形柱",
          etype="Column", qty=100.0),
    _rows(renovation="New", code="010505001", name="现浇混凝土梁",
          etype="Beam", qty=50.0),
    _rows(renovation="ToBeDemolished", code="010502001",
          name="现浇混凝土矩形柱", etype="Column", qty=30.0),
    _rows(renovation="Existing", code="010505003", name="现浇混凝土板",
          etype="Slab", qty=20.0),
]


@pytest.fixture()
def patch_estimate(monkeypatch):
    monkeypatch.setattr(
        instance_manager.manager, "get_info", lambda instance=None: _Info())
    calls: list[dict] = []

    def fake_tpl(name, instance=None, renovation_status=None,
                 split_by_renovation=False):
        calls.append({"name": name, "status": renovation_status,
                      "split": split_by_renovation})
        return {"template": name, "rows": list(_FLAT),
                "flat_rows": list(_FLAT), "warnings": []}

    monkeypatch.setattr(
        aggregate_service.service, "aggregate_by_template", fake_tpl)
    return calls


def test_estimate_project_split_groups_and_total(patch_estimate):
    out = _tool()(template_names=["concrete_volume"],
                  project_name="办公楼改造", split_by_renovation=True)
    assert out["status"] == "success"
    data = out["data"]
    assert data["mode"] == "split_by_renovation"
    assert patch_estimate[0]["split"] is True
    # 新建 2 条（柱/梁），拆除 1 条（柱），已建 1 条不计价
    assert data["new"]["item_count"] == 2
    assert data["demolition"]["item_count"] == 1
    assert any("已建 1 条不计价" in n for n in data["notes"])
    ns = data["new"]["summary"]["total"]
    ds = data["demolition"]["summary"]["total"]
    assert ds > 0 and ns > 0
    assert data["total"]["total"] == pytest.approx(round(ns + ds, 2))
    # 拆除组：人工费高于材料费（拆除更耗工，材料仅清运/损耗）
    dl = data["demolition"]["summary"]["labor_cost"]
    dm = data["demolition"]["summary"]["material_cost"]
    assert dl > dm
    assert data["demolition"]["summary"]["machine_cost"] > 0


def test_estimate_project_split_uses_own_consumption(patch_estimate):
    """同编码行：新建走建造定额，拆除走拆除定额，单价应不同。"""
    out = _tool()(template_names=["concrete_volume"],
                  project_name="P", split_by_renovation=True)
    data = out["data"]
    n_item = next(i for i in data["new"]["items"] if i["code"] == "010502001")
    d_item = next(i for i in data["demolition"]["items"]
                  if i["code"] == "010502001")
    # 单位价（拆除人工/机械更高）
    assert d_item["unit_price"] != n_item["unit_price"]
    assert d_item["labor_price"] / d_item["quantity"] > \
        n_item["labor_price"] / n_item["quantity"]


def test_estimate_project_single_demolition_status(patch_estimate):
    out = _tool()(template_names=["concrete_volume"],
                  project_name="办公楼改造", renovation_status="拆除")
    assert out["status"] == "success"
    data = out["data"]
    # 单表结构：无 new/demolition/total 分组
    assert "new" not in data and "summary" in data
    # 归一化后传给算量：拆除
    assert patch_estimate[0]["status"] == "ToBeDemolished"
    assert patch_estimate[0]["split"] is False
    # 选表=拆除定额 → 人工费高于材料费
    assert data["summary"]["labor_cost"] > data["summary"]["material_cost"]


def test_estimate_project_default_behavior_unchanged(patch_estimate):
    """不传新参数：返回结构与旧版一致（无 mode 键）。"""
    out = _tool()(template_names=["concrete_volume"], project_name="P")
    assert out["status"] == "success"
    assert set(out["data"]) == {
        "per_template", "summary", "coverage", "items", "notes",
        "fee_rates", "labor_price", "date"}
    assert patch_estimate[0]["split"] is False
    assert patch_estimate[0]["status"] is None


def test_estimate_project_keeps_multistory_rows_same_code(monkeypatch):
    """同编码多楼层行必须全部保留（量按 code 求和），不得只取第一行。"""
    monkeypatch.setattr(
        instance_manager.manager, "get_info", lambda instance=None: _Info())
    rows = [
        _rows(code="010502001", etype="Column", story=0, qty=10.0),
        _rows(code="010502001", etype="Column", story=1, qty=30.0),
        _rows(code="010502001", etype="Column", story=2, qty=5.0),
    ]

    def fake_tpl(name, instance=None, renovation_status=None,
                 split_by_renovation=False):
        return {"template": name, "rows": list(rows),
                "flat_rows": list(rows), "warnings": []}

    monkeypatch.setattr(
        aggregate_service.service, "aggregate_by_template", fake_tpl)
    out = _tool()(template_names=["column_concrete_volume"], project_name="P")
    item = out["data"]["items"][0]
    assert item["code"] == "010502001"
    assert item["quantity"] == pytest.approx(45.0)  # 10+30+5，非仅 10


def test_estimate_project_dedup_across_templates(monkeypatch):
    """两模板产出同编码：前者全部行保留，后者整块跳过。"""
    monkeypatch.setattr(
        instance_manager.manager, "get_info", lambda instance=None: _Info())
    a = [_rows(code="010502001", etype="Column", story=0, qty=10.0),
         _rows(code="010502001", etype="Column", story=1, qty=20.0)]
    b = [_rows(code="010502001", etype="Column", story=0, qty=999.0)]

    def fake_tpl(name, instance=None, renovation_status=None,
                 split_by_renovation=False):
        rows = a if name == "tpl_a" else b
        return {"template": name, "rows": list(rows),
                "flat_rows": list(rows), "warnings": []}

    monkeypatch.setattr(
        aggregate_service.service, "aggregate_by_template", fake_tpl)
    out = _tool()(template_names=["tpl_a", "tpl_b"], project_name="P")
    assert out["data"]["items"][0]["quantity"] == pytest.approx(30.0)
    assert out["data"]["per_template"]["tpl_b"]["kept"] == 0
    assert any("已由前序模板统计" in n for n in out["data"]["notes"])


# --- 多 sheet 导出 -------------------------------------------------------

def test_build_combined_estimate_report_three_sheets(tmp_path):
    new_table = estimate_from_flat_rows([_rows(renovation="New", qty=100.0)], "P")
    demo_table = estimate_from_flat_rows(
        [_rows(renovation="ToBeDemolished", qty=30.0)], "P",
        consumption_table=DEMOLITION_CONSUMPTION)
    total_table = combine_estimate_tables(new_table, demo_table, "P")
    path = build_combined_estimate_report(
        new_table, demo_table, total_table, "办公楼改造",
        "v59_combined", str(tmp_path))
    assert path.endswith(".xlsx")

    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert wb.sheetnames == ["新建概算", "拆除概算", "合计"]
    ws = wb["合计"]
    tr = next(r for r in range(1, ws.max_row + 1)
              if ws.cell(r, 1).value == "含税总价")
    assert ws.cell(tr, 7).value == pytest.approx(total_table["summary"]["total"])


def test_combined_report_theme_and_share(tmp_path):
    """V0.61：分状态表头着色 + 合计表追加「占比」列 + 冻结表头。"""
    new_table = estimate_from_flat_rows([_rows(renovation="New", qty=100.0)], "P")
    demo_table = estimate_from_flat_rows(
        [_rows(renovation="ToBeDemolished", qty=30.0)], "P",
        consumption_table=DEMOLITION_CONSUMPTION)
    total_table = combine_estimate_tables(new_table, demo_table, "P")
    path = build_combined_estimate_report(
        new_table, demo_table, total_table, "办公楼改造",
        "v61_theme", str(tmp_path))

    from openpyxl import load_workbook
    wb = load_workbook(path)

    # 各 sheet 表头底色区分（新建绿 / 拆除红 / 合计蓝灰）
    assert wb["新建概算"].cell(3, 1).fill.fgColor.rgb.endswith("C6EFCE")
    assert wb["拆除概算"].cell(3, 1).fill.fgColor.rgb.endswith("FFC7CE")
    assert wb["合计"].cell(3, 1).fill.fgColor.rgb.endswith("D9E1F2")

    # 新建/拆除表无「占比」列，合计表有
    assert "占比" not in [wb["新建概算"].cell(3, c).value for c in range(1, 12)]
    ws_total = wb["合计"]
    headers = [ws_total.cell(3, c).value
               for c in range(1, ws_total.max_column + 1)]
    assert "占比" in headers
    share_col = headers.index("占比") + 1
    denom = total_table["summary"]["total"]
    expected = round(new_table["summary"]["total"] / denom * 100, 1)
    assert ws_total.cell(4, share_col).value == f"{expected}%"

    # 冻结表头
    assert ws_total.freeze_panes == "A4"
