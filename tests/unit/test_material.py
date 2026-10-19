"""material_report 单测：单位标注 + 翻新分组 + 材料费联动 + Excel 导出（V0.62）。"""

from archilinkai import i18n
from archilinkai import material_report
from archilinkai.material_report import (
    MATERIAL_UNITS,
    _material_unit_price,
    export_material_excel,
    generate_material_report,
)


class _Page:
    def __init__(self, items):
        self.items = items


def _setup(monkeypatch, elems, props_map):
    """mock engine.query 与 property_service.service.get_values。"""
    monkeypatch.setattr(material_report.engine, "query",
                        lambda req: _Page(elems))

    class FakeProp:
        def get_values(self, guids, names, instance=None):
            return {g: {n: props_map.get(g, {}).get(n) for n in names}
                    for g in guids}

    monkeypatch.setattr(material_report.property_service, "service", FakeProp())


def test_column_units_and_cost(monkeypatch):
    elems = [{"guid": "g1", "name": "KZ1", "story_index": 0}]
    _setup(monkeypatch, elems, {"g1": {"General_Volume": 100.0}})
    r = generate_material_report(["Column"], "inst")

    # 单位标注
    assert r["units"]["混凝土"] == "m³"
    assert r["units"]["钢筋"] == "kg"
    assert r["units"]["模板"] == "m²"
    # 用量（Column 因子：混凝土 1.0 / 钢筋 150 / 模板 7.2）
    assert r["summary"]["混凝土"] == 100.0
    assert r["summary"]["钢筋"] == 15000.0
    assert r["summary"]["模板"] == 720.0
    # 材料费（混凝土 380/m³、钢筋 4/kg、模板 45/m²）
    assert r["price"]["混凝土"] == {"unit_price": 380.0, "cost": 38000.0}
    assert r["price"]["钢筋"] == {"unit_price": 4.0, "cost": 60000.0}
    assert r["price"]["模板"] == {"unit_price": 45.0, "cost": 32400.0}
    assert r["material_cost_total"] == 130400.0


def test_wall_volume_to_area_and_unpriced(monkeypatch):
    # Wall 20 m³、墙厚 0.2 → 面积 100 m²
    elems = [{"guid": "g1", "name": "Q1", "story_index": 0}]
    _setup(monkeypatch, elems, {"g1": {
        "General_NetVolume": 20.0, "General_Thickness": 0.2}})
    r = generate_material_report(["Wall"], "inst")

    assert r["summary"]["砖砌体"] == 20.0
    assert r["summary"]["内墙涂料"] == 200.0   # 双面
    assert r["summary"]["外墙涂料"] == 100.0   # 单面
    assert r["summary"]["保温板"] == 100.0
    assert r["summary"]["防水涂料"] == 20.0
    # 防水涂料无参考价 → 不计价
    assert r["price"]["防水涂料"] == {"unit_price": None, "cost": None}
    # 合计只累加已定价的材料
    assert r["material_cost_total"] > 0


def test_split_by_renovation(monkeypatch):
    elems = [
        {"guid": "g1", "name": "A", "story_index": 0},
        {"guid": "g2", "name": "B", "story_index": 1},
    ]
    props = {
        "g1": {"General_Volume": 10.0, "Category_RenovationStatus": "New"},
        "g2": {"General_Volume": 30.0,
               "Category_RenovationStatus": "ToBeDemolished"},
    }
    _setup(monkeypatch, elems, props)
    r = generate_material_report(["Slab"], "inst", split_by_renovation=True)

    assert "by_renovation" in r
    assert r["by_renovation"]["New"]["混凝土"] == 10.0
    assert r["by_renovation"]["ToBeDemolished"]["混凝土"] == 30.0
    # 明细带翻新状态
    revs = {d["guid"]: d["renovation"] for d in r["details"]}
    assert revs == {"g1": "New", "g2": "ToBeDemolished"}


def test_no_split_when_disabled(monkeypatch):
    elems = [{"guid": "g1", "name": "A", "story_index": 0}]
    _setup(monkeypatch, elems, {"g1": {"General_Volume": 5.0}})
    r = generate_material_report(["Slab"], "inst")
    assert "by_renovation" not in r
    assert "renovation" not in r["details"][0]


def test_material_unit_price_lookup():
    assert _material_unit_price("混凝土") == 380.0
    assert _material_unit_price("钢筋") == 4.0        # 4000 元/t × 0.001
    assert _material_unit_price("保温板") == 19.0     # 380 元/m³ × 0.05
    assert _material_unit_price("不存在") is None
    assert MATERIAL_UNITS["钢筋"] == "kg"


def test_export_material_excel(monkeypatch, tmp_path):
    elems = [
        {"guid": "g1", "name": "A", "story_index": 0},
        {"guid": "g2", "name": "B", "story_index": 1},
    ]
    props = {
        "g1": {"General_Volume": 10.0, "Category_RenovationStatus": "New"},
        "g2": {"General_Volume": 20.0, "Category_RenovationStatus": "New"},
    }
    _setup(monkeypatch, elems, props)
    r = generate_material_report(["Column"], "inst", split_by_renovation=True)

    path = str(tmp_path / "m.xlsx")
    assert export_material_excel(r, path) == path

    from openpyxl import load_workbook
    wb = load_workbook(path)
    summary = i18n.t("material.sheet_summary", "材料汇总")
    detail = i18n.t("material.sheet_detail", "明细")
    assert summary in wb.sheetnames
    assert detail in wb.sheetnames
    assert "Column" in wb.sheetnames
    # 含「按翻新状态」分表
    assert any(s.startswith(i18n.t("material.sheet_renovation", "按翻新状态"))
               for s in wb.sheetnames)
    # 汇总页表头有单位/单价/材料费列 + 表头着色 + 冻结
    ws = wb[summary]
    headers = [ws.cell(1, c).value for c in range(1, 6)]
    assert i18n.t("material.h_unit", "单位") in headers
    assert i18n.t("material.h_cost", "材料费(元)") in headers
    assert ws.cell(1, 1).fill.fgColor.rgb.endswith("D9E1F2")
    assert ws.freeze_panes == "A2"
    # 明细页含翻新状态列
    ws_d = wb[detail]
    assert i18n.t("material.h_renovation", "翻新状态") in [
        ws_d.cell(1, c).value for c in range(1, 8)]
